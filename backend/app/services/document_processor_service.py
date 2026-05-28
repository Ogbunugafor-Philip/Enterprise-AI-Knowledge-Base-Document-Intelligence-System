import re
import unicodedata
from pathlib import Path


def extract_text_from_pdf(file_path: str | Path) -> tuple[str, int]:
    try:
        import fitz  # type: ignore

        doc = fitz.open(file_path)
        if doc.is_encrypted:
            return "", 0
        return "\n".join(page.get_text() for page in doc), len(doc)
    except Exception:
        try:
            import pdfplumber  # type: ignore

            with pdfplumber.open(file_path) as pdf:
                return "\n".join(page.extract_text() or "" for page in pdf.pages), len(pdf.pages)
        except Exception:
            return "", 0


def extract_text_from_docx(file_path: str | Path) -> str:
    try:
        import docx  # type: ignore

        document = docx.Document(file_path)
        return "\n\n".join(paragraph.text for paragraph in document.paragraphs)
    except Exception:
        return ""


def extract_text_from_txt(file_path: str | Path) -> str:
    data = Path(file_path).read_bytes()
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def extract_text_from_excel(file_path: str | Path) -> str:
    try:
        import openpyxl  # type: ignore

        workbook = openpyxl.load_workbook(file_path, data_only=True)
        rows = []
        for sheet in workbook.worksheets:
            rows.append(f"# {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                rows.append(" | ".join("" if cell is None else str(cell) for cell in row))
        return "\n".join(rows)
    except Exception:
        return ""


def extract_text_from_image(file_path: str | Path) -> tuple[str, float]:
    try:
        import pytesseract  # type: ignore
        from PIL import Image  # type: ignore

        return pytesseract.image_to_string(Image.open(file_path)), 0.8
    except Exception:
        return "", 0.0


def route_extraction(file_path: str | Path, file_type: str) -> dict:
    extension = file_type.lower().lstrip(".")
    if extension == "pdf":
        text, pages = extract_text_from_pdf(file_path)
        return {"text": text, "page_count": pages, "method": "pdf"}
    if extension == "docx":
        return {"text": extract_text_from_docx(file_path), "page_count": None, "method": "docx"}
    if extension == "txt":
        return {"text": extract_text_from_txt(file_path), "page_count": None, "method": "txt"}
    if extension in {"xlsx", "xls"}:
        return {"text": extract_text_from_excel(file_path), "page_count": None, "method": "excel"}
    if extension in {"png", "jpg", "jpeg", "tiff"}:
        text, confidence = extract_text_from_image(file_path)
        return {"text": text, "page_count": None, "method": "ocr", "confidence": confidence}
    return {"text": "", "page_count": None, "method": "unsupported"}


def clean_extracted_text(text: str) -> str:
    text = text.replace("\\x00", "")
    normalized = unicodedata.normalize("NFKC", text.replace("\x00", ""))
    normalized = re.sub(r"[ \t]+", " ", normalized)
    normalized = re.sub(r"(\\n){3,}", r"\\n\\n", normalized)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    return normalized.strip()


def preprocess_for_chunking(text: str) -> dict:
    sections = []
    current = {"heading": "Document", "paragraphs": []}
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue
        if len(stripped) <= 120 and (stripped.isupper() or stripped.startswith("#")):
            if current["paragraphs"]:
                sections.append(current)
            current = {"heading": stripped.lstrip("# "), "paragraphs": []}
        else:
            current["paragraphs"].append(stripped)
    if current["paragraphs"]:
        sections.append(current)
    return {"sections": sections, "text": text}
